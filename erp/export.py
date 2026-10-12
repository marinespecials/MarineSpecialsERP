"""Excel exports (openpyxl only - no pandas needed)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="0B3C5D")
EURO = '#,##0.00 "€"'

LEDGER_COLUMNS = [
    ("Invoice Date", "invoice_date", 12), ("Supplier", "supplier_name", 28), ("ΑΦΜ", "supplier_vat", 12),
    ("Invoice #", "invoice_no", 16), ("Net", "net", 12), ("VAT", "vat_amount", 12), ("Total", "total", 13),
    ("Due", "due_date", 12), ("Status", "status", 10), ("Paid On", "paid_date", 12),
    ("ΜΑΡΚ", "mark", 18), ("Stock", "stock_status", 9), ("File", "file_path", 50),
]


def _sheet(ws, columns, rows):
    for c, (title, _, width) in enumerate(columns, 1):
        cell = ws.cell(1, c, title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = width
    for r, row in enumerate(rows, 2):
        for c, (title, key, _) in enumerate(columns, 1):
            v = row.get(key)
            cell = ws.cell(r, c, v if v not in ("", None) else None)
            if title in ("Net", "VAT", "Total", "Spend"):
                cell.number_format = EURO
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def export_ledger(path, invoices, monthly=None, suppliers=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Ledger"
    _sheet(ws, LEDGER_COLUMNS, invoices)
    if monthly is not None:
        _sheet(wb.create_sheet("By Month"), [("Month", "ym", 10), ("Invoices", "n", 10), ("Spend", "spend", 14)], monthly)
    if suppliers is not None:
        _sheet(wb.create_sheet("By Supplier"),
               [("Supplier", "supplier_name", 30), ("Invoices", "n", 10), ("Spend", "spend", 14)], suppliers)
    wb.save(path)
    return path
